#!/usr/bin/env python3
"""Genera catálogos TypeScript limpios desde las exportaciones de VINTO."""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


MACHINE_ALIASES = {
    "SINCRO 2": "Sincro II",
    "SINCRO 3": "Sincro III",
    "OMET 1": "OMET I",
    "OMET 2": "OMET II",
    "SERVILLETERA 4": "Serv 4",
}


def text(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def code(value: object) -> str:
    return text(value).upper()


def machine(value: object) -> str:
    raw = text(value).upper()
    return MACHINE_ALIASES.get(raw, raw if raw in {"MP1", "MP3"} else text(value))


def read_inventory(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["INInventario"]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    index = {name: position for position, name in enumerate(headers)}
    inventory = {}
    for row in rows:
        article = row[index["Articulo"]]
        if article is None:
            continue
        inventory[code(article)] = {
            "code": code(article),
            "name": text(row[index["Descripcion"]]),
            "className": text(row[index["ClaseArticulo"]]),
            "unit": text(row[index["UM"]]),
        }
    return workbook, inventory


def read_materials(workbook, inventory):
    sheet = workbook["MFRutaMaterialOP"]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    index = {name: position for position, name in enumerate(headers)}
    materials = {}
    for row in rows:
        raw_code = row[index["Material"]]
        if raw_code is None:
            continue
        material_code = code(raw_code)
        official = inventory.get(material_code)
        materials[material_code] = {
            "code": material_code,
            "name": official["name"] if official else text(row[index["Descripcion"]]),
            "className": official["className"] if official else "SIN CLASIFICAR",
            "unit": official["unit"] if official else text(row[index["UM"]]),
        }
    return sorted(materials.values(), key=lambda item: (item["className"], item["name"], item["code"]))


def read_machine_products(path: Path, inventory):
    workbook = load_workbook(path, read_only=True, data_only=True)
    pairs = {}
    for sheet_name in ("BOBINAS", "CONV"):
        sheet = workbook[sheet_name]
        current_machine = None
        for row in sheet.iter_rows(min_row=4, values_only=True):
            if row[0] is not None:
                current_machine = machine(row[0])
            if row[1] is None or current_machine is None:
                continue
            product_code = code(row[1])
            official = inventory[product_code]
            pairs[(current_machine, product_code)] = {
                "machine": current_machine,
                "code": product_code,
                "name": official["name"],
                "unit": official["unit"],
            }

    reb_sheet = workbook["Reb"]
    for row in reb_sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None or row[1] is None:
            continue
        current_machine = machine(row[0])
        product_code = code(row[1])
        official = inventory[product_code]
        pairs[(current_machine, product_code)] = {
            "machine": current_machine,
            "code": product_code,
            "name": official["name"],
            "unit": official["unit"],
        }

    grouped = defaultdict(list)
    for item in pairs.values():
        grouped[item["machine"]].append({key: value for key, value in item.items() if key != "machine"})
    return {
        key: sorted(value, key=lambda item: (item["name"], item["code"]))
        for key, value in sorted(grouped.items())
    }


def ts_export(name: str, value: object) -> str:
    return f"export const {name} = {json.dumps(value, ensure_ascii=False, indent=2)} as const;\n"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--inventory", required=True, type=Path)
    parser.add_argument("--machines", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    workbook, inventory = read_inventory(args.inventory)
    materials = read_materials(workbook, inventory)
    products = read_machine_products(args.machines, inventory)

    content = "// Archivo generado desde los catálogos oficiales entregados por Data VINTO.\n"
    content += "// Los códigos se normalizan a mayúsculas y las relaciones duplicadas se eliminan.\n\n"
    content += "export type CatalogItem = { code: string; name: string; unit: string };\n"
    content += "export type MaterialCatalogItem = CatalogItem & { className: string };\n\n"
    content += ts_export("PRODUCTS_BY_MACHINE", products)
    content += "\n"
    content += ts_export("MATERIALS", materials)
    args.output.write_text(content, encoding="utf-8")


if __name__ == "__main__":
    main()
